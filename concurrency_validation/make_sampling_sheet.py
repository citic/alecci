#!/usr/bin/env python3
"""Genera la hoja de cálculo de muestreo aleatorio para la revisión manual.

Una hoja por dataset. Cada fila es un ejercicio (un archivo .ale). A cada
ejercicio se le asigna un número pseudoaleatorio en [0, 1) con una semilla fija
(reproducible), se ordenan las filas por ese número y se marca el primer
PORCENTAJE % (redondeado hacia arriba) para revisión.

Uso:
    python3 make_sampling_sheet.py [--seed 42] [--pct 5] [--out muestreo.xlsx]

Requiere: openpyxl.
"""
import argparse
import math
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent / "alecci_benchmarks"
DATASETS = {
    "ClassExamples": ROOT / "labeled_benchmarks" / "ClassExamples",
    "DataRaceBench": ROOT / "labeled_benchmarks" / "DataRaceBench",
    "DeepRace": ROOT / "labeled_benchmarks" / "DeepRace",
    "SCTBench": ROOT / "SCTBench",
}


def list_exercises(base: Path):
    # Orden determinista: la reproducibilidad depende de él.
    return sorted(p.relative_to(base).as_posix() for p in base.rglob("*.ale"))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--seed", type=int, default=42)
    ap.add_argument("--pct", type=float, default=5.0)
    ap.add_argument("--out", default=str(Path(__file__).resolve().parent / "muestreo.xlsx"))
    args = ap.parse_args()

    wb = Workbook()
    meta = wb.active
    meta.title = "Método"
    bold = Font(bold=True)
    hl = PatternFill("solid", fgColor="FFF2CC")

    meta_rows = [
        ("Semilla (random.Random)", args.seed),
        ("Porcentaje a revisar", args.pct / 100),
        ("Generador", "Python random.Random(seed) (Mersenne Twister), un generador por dataset: Random(f'{seed}-{dataset}')"),
        ("Unidad", "Cada archivo .ale es un ejercicio"),
        ("Procedimiento", "1) listar ejercicios en orden alfabético; 2) asignar random() a cada uno; "
                          "3) ordenar por ese número; 4) revisar los primeros ceil(N × porcentaje)"),
        ("Reproducir", "python3 concurrency_validation/make_sampling_sheet.py --seed %d --pct %g" % (args.seed, args.pct)),
        ("", ""),
        ("Dataset", "Total ejercicios / a revisar"),
    ]
    for r in meta_rows:
        meta.append(r)
    meta["A1"].font = meta["A2"].font = bold
    meta["B2"].number_format = "0%"
    meta["B2"].fill = hl  # editable: cambia el % y la columna Revisar se recalcula
    meta["A8"].font = meta["B8"].font = bold

    for name, base in DATASETS.items():
        files = list_exercises(base)
        rng = random.Random(f"{args.seed}-{name}")
        rows = [(f, rng.random()) for f in files]
        rows.sort(key=lambda t: t[1])
        n = len(rows)

        ws = wb.create_sheet(name)
        ws.append(["Posición", "Ejercicio", "Aleatorio [0,1)", "Revisar", "Revisado", "Notas"])
        for c in ws[1]:
            c.font = bold
        for i, (f, x) in enumerate(rows, start=1):
            r = i + 1
            ws.append([i, f, x, f'=IF(A{r}<=ROUNDUP(COUNTA($B$2:$B${n + 1})*Método!$B$2,0),"SÍ","")', "", ""])
        k = math.ceil(n * args.pct / 100)
        for r in range(2, k + 2):
            for c in ws[r]:
                c.fill = hl
        ws.freeze_panes = "A2"
        for col, w in zip("ABCDEF", (10, 70, 18, 10, 10, 40)):
            ws.column_dimensions[col].width = w
        meta.append((name, f"{n} / {k}"))

    meta.column_dimensions["A"].width = 28
    meta.column_dimensions["B"].width = 110

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"Escrito {out}")


if __name__ == "__main__":
    main()
